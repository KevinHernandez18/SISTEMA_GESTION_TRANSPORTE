import os
import sys
import django
import io
# Ensure project root is on sys.path so Django settings can be imported
project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if project_root not in sys.path:
    sys.path.insert(0, project_root)
os.environ.setdefault('DJANGO_SETTINGS_MODULE','backend.settings')
django.setup()
from django.test import Client
from rest_framework_simplejwt.tokens import RefreshToken
from django.contrib.auth import get_user_model
User=get_user_model()
user=User.objects.filter(username='ordering_test_user').first()
if not user:
    user=User.objects.create_superuser('ordering_test_user','test@example.com','TestPassword123!')
access=str(RefreshToken.for_user(user).access_token)
client=Client()
# Specify a valid HTTP_HOST to avoid DisallowedHost in tests
res=client.get('/gestion_transporte/api/v1/licencias/export/', HTTP_AUTHORIZATION=f'Bearer {access}', HTTP_HOST='127.0.0.1:8000')
from openpyxl import load_workbook
wb=load_workbook(io.BytesIO(res.content))
ws=wb.active
headers=[cell.value for cell in next(ws.iter_rows(min_row=1,max_row=1))]
print('STATUS',res.status_code)
print('HEADERS',headers)
